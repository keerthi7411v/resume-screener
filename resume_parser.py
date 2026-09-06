"""Extract raw text from PDF and DOCX resumes."""
import os
from PyPDF2 import PdfReader
from docx import Document


def extract_from_pdf(path):
    reader = PdfReader(path)
    text = []
    for page in reader.pages:
        text.append(page.extract_text() or "")
    return "\n".join(text)


def extract_from_docx(path):
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:                     # some resumes use tables
        for row in table.rows:
            parts.append(" ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def extract_text(path):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        return extract_from_pdf(path)
    if ext == ".docx":
        return extract_from_docx(path)
    if ext == ".txt":
        with open(path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()
    raise ValueError(f"Unsupported file type: {ext}")
